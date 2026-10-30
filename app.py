import streamlit as st
import collections
import collections.abc
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
import os
import io
import re
import posixpath
import tempfile
import warnings
import base64
import colorsys
import hashlib
from PIL import Image
from collections import Counter
from copy import deepcopy
from lxml import etree
import pptx
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI

# Critical patch for python-pptx: templates containing SVGs lack sha1 on generic Part,
# which causes get_or_add_image_part() / add_picture() to crash with AttributeError
if not hasattr(Part, 'sha1'):
    Part.sha1 = property(lambda self: hashlib.sha1(self.blob).hexdigest() if hasattr(self, 'blob') else '')

warnings.filterwarnings('ignore', category=UserWarning, module='zipfile')

# Set Streamlit Page Config
st.set_page_config(
    page_title="IfG SlideStyler | PowerPoint Vorlagen-Engine",
    page_icon="🏢",
    layout="centered",
    initial_sidebar_state="collapsed"
)

# Load IfG Logo as base64 for reliable inline rendering
logo_b64 = ""
if os.path.exists("ifg-logo.png"):
    with open("ifg-logo.png", "rb") as f:
        logo_b64 = base64.b64encode(f.read()).decode("utf-8")

# Custom CSS matching https://qr.ifg-zukunftswerkstatt.de/
st.markdown("""
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&display=swap');

    /* Global Background & Typography */
    html, body, [data-testid="stAppViewContainer"], .main {
        background-color: #f8fafc !important;
        color: #0f172a !important;
        font-family: 'Inter', -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif !important;
    }

    /* Hide Streamlit Header, GitHub ribbon & Watermarks */
    #MainMenu {visibility: hidden !important; display: none !important;}
    header {visibility: hidden !important; display: none !important;}
    footer {visibility: hidden !important; display: none !important;}
    [data-testid="stHeader"] {display: none !important;}
    [data-testid="stToolbar"] {visibility: hidden !important; display: none !important;}
    [data-testid="stDecoration"] {display: none !important;}
    [data-testid="stStatusWidget"] {display: none !important;}
    .stAppDeployButton {display: none !important;}

    .block-container {
        padding-top: 1.5rem !important;
        padding-bottom: 3.5rem !important;
        max-width: 860px !important;
    }

    /* Typography */
    h1, h2, h3, h4, h5, h6 {
        color: #0f172a !important;
        font-family: 'Inter', sans-serif !important;
        font-weight: 700 !important;
        letter-spacing: -0.025em !important;
    }

    p, span, label, li {
        color: #334155;
    }

    /* Top Brand Navigation Header */
    .ifg-top-nav {
        background: #ffffff;
        border: 1px solid #e2e8f0;
        border-radius: 16px;
        box-shadow: 0 1px 3px 0 rgba(0, 0, 0, 0.05);
        padding: 14px 22px;
        margin-bottom: 24px;
        display: flex;
        align-items: center;
        justify-content: space-between;
        flex-wrap: wrap;
        gap: 12px;
    }

    .ifg-nav-left {
        display: flex;
        align-items: center;
        gap: 16px;
    }

    .ifg-logo-img {
        height: 38px;
        width: auto;
        object-fit: contain;
    }

    .ifg-nav-divider {
        height: 32px;
        width: 1px;
        background-color: #e2e8f0;
    }

    .ifg-nav-title {
        font-size: 1.18rem;
        font-weight: 700;
        color: #0f172a;
        margin: 0;
        display: flex;
        align-items: center;
        gap: 8px;
    }

    .ifg-badge-tool {
        display: inline-flex;
        align-items: center;
        padding: 2px 8px;
        border-radius: 6px;
        font-size: 0.72rem;
        font-weight: 600;
        background-color: #eff6ff;
        color: #012641;
        border: 1px solid #bfdbfe;
    }

    .ifg-nav-subtitle {
        font-size: 0.8rem;
        color: #64748b;
        margin: 2px 0 0 0;
    }

    .ifg-nav-badge {
        display: inline-flex;
        align-items: center;
        gap: 6px;
        font-size: 0.75rem;
        font-weight: 500;
        color: #047857;
        background-color: #ecfdf5;
        border: 1px solid #a7f3d0;
        padding: 5px 12px;
        border-radius: 9999px;
    }

    /* Cards */
    .ifg-card {
        background: #ffffff;
        border: 1px solid #e2e8f0;
        border-radius: 16px;
        box-shadow: 0 1px 3px 0 rgba(0, 0, 0, 0.05);
        padding: 22px 24px;
        margin-bottom: 22px;
    }

    .ifg-card-header {
        display: flex;
        align-items: center;
        gap: 10px;
        padding-bottom: 12px;
        border-bottom: 1px solid #f1f5f9;
        margin-bottom: 14px;
    }

    .ifg-card-title {
        font-size: 1.05rem;
        font-weight: 700;
        color: #012641;
        margin: 0;
    }

    .ifg-feature-list {
        list-style: none;
        padding-left: 0;
        margin: 0;
    }

    .ifg-feature-item {
        display: flex;
        align-items: flex-start;
        gap: 10px;
        padding: 6px 0;
        font-size: 0.92rem;
        color: #334155;
        line-height: 1.5;
    }

    .ifg-bullet-dot {
        width: 8px;
        height: 8px;
        border-radius: 50%;
        background-color: #3EB9C0;
        margin-top: 7px;
        flex-shrink: 0;
    }

    /* Streamlit Form Components */
    [data-testid="stSelectbox"] > div > div {
        background-color: #ffffff !important;
        border: 1px solid #cbd5e1 !important;
        border-radius: 10px !important;
        color: #0f172a !important;
        box-shadow: 0 1px 2px 0 rgba(0, 0, 0, 0.04) !important;
        font-size: 0.92rem !important;
    }
    [data-testid="stSelectbox"] > div > div:hover {
        border-color: #012641 !important;
    }

    [data-testid="stFileUploader"] section {
        background-color: #ffffff !important;
        border: 2px dashed #cbd5e1 !important;
        border-radius: 14px !important;
        padding: 20px !important;
        transition: all 0.2s ease !important;
    }
    [data-testid="stFileUploader"] section:hover {
        border-color: #012641 !important;
        background-color: #f8fafc !important;
    }

    label, [data-testid="stWidgetLabel"] p {
        color: #0f172a !important;
        font-weight: 600 !important;
        font-size: 0.92rem !important;
    }

    /* Buttons */
    .stButton > button {
        background: #012641 !important;
        color: #ffffff !important;
        border: none !important;
        border-radius: 12px !important;
        padding: 0.85rem 2rem !important;
        font-size: 1.05rem !important;
        font-weight: 600 !important;
        width: 100% !important;
        box-shadow: 0 2px 4px 0 rgba(1, 38, 65, 0.18) !important;
        transition: all 0.2s ease !important;
    }
    .stButton > button:hover {
        background: #083254 !important;
        transform: translateY(-1px);
        box-shadow: 0 4px 10px 0 rgba(1, 38, 65, 0.25) !important;
    }

    .stDownloadButton > button {
        background: #059669 !important;
        color: #ffffff !important;
        border: none !important;
        border-radius: 12px !important;
        padding: 0.85rem 2rem !important;
        font-size: 1.05rem !important;
        font-weight: 600 !important;
        width: 100% !important;
        box-shadow: 0 2px 4px 0 rgba(5, 150, 105, 0.2) !important;
        transition: all 0.2s ease !important;
    }
    .stDownloadButton > button:hover {
        background: #047857 !important;
        transform: translateY(-1px);
        box-shadow: 0 4px 10px 0 rgba(5, 150, 105, 0.28) !important;
    }

    /* Success Card */
    .ifg-success-box {
        background: #f0fdf4;
        border: 1px solid #bbf7d0;
        border-radius: 14px;
        padding: 20px;
        margin-top: 20px;
        margin-bottom: 20px;
        text-align: center;
    }
    .ifg-success-title {
        font-size: 1.15rem;
        font-weight: 700;
        color: #15803d;
        margin-bottom: 6px;
    }
    .ifg-success-desc {
        font-size: 0.92rem;
        color: #166534;
        margin: 0;
    }

    /* Expander Styling */
    [data-testid="stExpander"] {
        background: #ffffff !important;
        border: 1px solid #e2e8f0 !important;
        border-radius: 12px !important;
        box-shadow: 0 1px 2px 0 rgba(0, 0, 0, 0.03) !important;
    }
    </style>
""", unsafe_allow_html=True)

# ----------------- OPENXML NAMESPACES & CORPORATE CONSTANTS -----------------

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
C_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
DGM_NS = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
DSP_NS = "http://schemas.microsoft.com/office/drawing/2008/diagram"

CORPORATE_FONT = "Abadi"

# Official IfG / Arsipa Corporate Color Palette (from Vorlage theme1.xml)
IFG_NAVY = "012641"       # Primary Brand Dark Navy (Headings, Primary Boxes, Headers)
IFG_CHARCOAL = "3A4A63"   # Body Text & Neutral Lines
IFG_MUTED = "8C96A5"      # Kickers, Subtitles, Muted Captions
IFG_WHITE = "FFFFFF"

DARK_BRAND_SWATCHES = ["012641", "006174", "3E3D5E", "007A8C"]
MID_BRAND_SWATCHES = ["00A8B7", "3EB9C0", "57577C", "2B7A9B"]
LIGHT_BRAND_SWATCHES = ["D9F0F2", "C5E4EE", "DCE3F0", "E8F5F6"]

SYMBOL_FONTS = {
    "wingdings", "wingdings 2", "wingdings 3", "webdings", "symbol",
    "marlett", "ms outlook", "bookshelf symbol 7", "segoe mdl2 assets",
    "segoe fluent icons", "fontawesome", "font awesome 5 free",
    "font awesome 6 free", "material icons"
}

PRESET_COLORS = {
    "black": (0, 0, 0),
    "white": (255, 255, 255),
    "red": (255, 0, 0),
    "green": (0, 128, 0),
    "blue": (0, 0, 255),
    "yellow": (255, 255, 0),
    "gray": (128, 128, 128),
    "grey": (128, 128, 128),
    "lightgray": (211, 211, 211),
    "darkgray": (105, 105, 105),
    "navy": (0, 0, 128),
    "teal": (0, 128, 128),
    "orange": (255, 165, 0),
}


def qn_a(tag):
    return f"{{{A_NS}}}{tag}"


def qn_p(tag):
    return f"{{{P_NS}}}{tag}"


def safe_shape_type(shape):
    """Safely retrieves shape_type without raising NotImplementedError on freeforms or exotic shapes."""
    try:
        return shape.shape_type
    except Exception:
        return None


def scan_arsipa_template(template_prs):
    """Finds the clean corporate canvas layout ('Leere Folie') in the Vorlage."""
    canvas_idx = 0
    for idx, layout in enumerate(template_prs.slide_layouts):
        name_lower = layout.name.lower() if layout.name else ""
        if "leere folie" in name_lower or "blank" in name_lower:
            canvas_idx = idx
            break
    return {"canvas_leere_folie": canvas_idx}


# ----------------- 1. COLOR & THEME RESOLUTION ENGINE -----------------

def hex_to_rgb(hex_str):
    h = hex_str.strip().lstrip("#")
    if len(h) != 6:
        return (0, 0, 0)
    return (int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def rgb_to_hex(rgb):
    return "{:02X}{:02X}{:02X}".format(
        max(0, min(255, int(round(rgb[0])))),
        max(0, min(255, int(round(rgb[1])))),
        max(0, min(255, int(round(rgb[2]))))
    )


def perceived_brightness(rgb):
    """Perceived brightness in [0.0, 1.0] matching human contrast perception."""
    return (0.299 * rgb[0] + 0.587 * rgb[1] + 0.114 * rgb[2]) / 255.0


def rgb_saturation(rgb):
    _, s, _ = colorsys.rgb_to_hsv(rgb[0] / 255.0, rgb[1] / 255.0, rgb[2] / 255.0)
    return s


def rgb_distance(c1, c2):
    return ((c1[0] - c2[0]) ** 2 + (c1[1] - c2[1]) ** 2 + (c1[2] - c2[2]) ** 2) ** 0.5


def shift_hex_lightness(hex_str, delta_l):
    r, g, b = hex_to_rgb(hex_str)
    h, l, s = colorsys.rgb_to_hls(r / 255.0, g / 255.0, b / 255.0)
    l_new = max(0.08, min(0.95, l + delta_l))
    rn, gn, bn = colorsys.hls_to_rgb(h, l_new, s)
    return rgb_to_hex((rn * 255.0, gn * 255.0, bn * 255.0))


def apply_oxml_color_mods(base_rgb, clr_el):
    """Applies OpenXML DrawingML color transforms (lumMod, lumOff, tint, shade, satMod, alpha)."""
    r, g, b = [c / 255.0 for c in base_rgb]
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    alpha_val = None

    for ch in clr_el:
        local = ch.tag.split("}")[-1]
        raw_val = ch.get("val")
        if raw_val is None:
            continue
        try:
            v = int(raw_val) / 100000.0
        except ValueError:
            continue

        if local == "alpha":
            alpha_val = ch.get("val")
        elif local == "lumMod":
            l = max(0.0, min(1.0, l * v))
        elif local == "lumOff":
            l = max(0.0, min(1.0, l + v))
        elif local == "satMod":
            s = max(0.0, min(1.0, s * v))
        elif local == "satOff":
            s = max(0.0, min(1.0, s + v))
        elif local == "tint":
            r, g, b = colorsys.hls_to_rgb(h, l, s)
            r = 1.0 - (1.0 - r) * v
            g = 1.0 - (1.0 - g) * v
            b = 1.0 - (1.0 - b) * v
            h, l, s = colorsys.rgb_to_hls(r, g, b)
        elif local == "shade":
            r, g, b = colorsys.hls_to_rgb(h, l, s)
            r = r * v
            g = g * v
            b = b * v
            h, l, s = colorsys.rgb_to_hls(r, g, b)

    rf, gf, bf = colorsys.hls_to_rgb(h, l, s)
    return (int(round(rf * 255)), int(round(gf * 255)), int(round(bf * 255))), alpha_val


def extract_slide_theme_map(slide):
    """Resolves the source slide's active theme color scheme + clrMap aliases into RGB tuples."""
    theme_map = {
        "dk1": (0, 0, 0),
        "lt1": (255, 255, 255),
        "dk2": (68, 84, 106),
        "lt2": (231, 230, 230),
        "accent1": (68, 114, 196),
        "accent2": (237, 125, 49),
        "accent3": (165, 165, 165),
        "accent4": (255, 192, 0),
        "accent5": (91, 155, 213),
        "accent6": (112, 173, 71),
        "hlink": (5, 99, 193),
        "folHlink": (149, 79, 114),
    }

    try:
        master = slide.slide_layout.slide_master
        for rel in master.part.rels.values():
            if not rel.is_external and rel.reltype.endswith("/theme"):
                t_root = etree.fromstring(rel.target_part.blob)
                clr_scheme = t_root.find(f".//{{{A_NS}}}clrScheme")
                if clr_scheme is not None:
                    for slot in clr_scheme:
                        slot_name = slot.tag.split("}")[-1]
                        for c in slot:
                            ctag = c.tag.split("}")[-1]
                            if ctag == "srgbClr" and c.get("val"):
                                theme_map[slot_name] = hex_to_rgb(c.get("val"))
                            elif ctag == "sysClr":
                                last = c.get("lastClr") or "000000"
                                theme_map[slot_name] = hex_to_rgb(last)
                break

        # Apply master clrMap (e.g. bg1->lt1, tx1->dk1, bg2->lt2, tx2->dk2)
        clr_map_el = master._element.find(f"{{{P_NS}}}clrMap")
        aliases = {"bg1": "lt1", "tx1": "dk1", "bg2": "lt2", "tx2": "dk2"}
        if clr_map_el is not None:
            for k, v in clr_map_el.attrib.items():
                aliases[k] = v

        # Check layout & slide clrMapOvr if overridden
        for obj in (slide.slide_layout._element, slide._element):
            ovr = obj.find(f"{{{P_NS}}}clrMapOvr/{{{A_NS}}}overrideClrMapping")
            if ovr is not None:
                for k, v in ovr.attrib.items():
                    aliases[k] = v

        for alias_k, target_k in aliases.items():
            if target_k in theme_map:
                theme_map[alias_k] = theme_map[target_k]
    except Exception:
        pass

    if "tx1" not in theme_map:
        theme_map["tx1"] = theme_map["dk1"]
    if "bg1" not in theme_map:
        theme_map["bg1"] = theme_map["lt1"]
    if "tx2" not in theme_map:
        theme_map["tx2"] = theme_map["dk2"]
    if "bg2" not in theme_map:
        theme_map["bg2"] = theme_map["lt2"]

    return theme_map


def resolve_color_from_parent(parent_el, theme_map, ph_rgb=None):
    """Finds and resolves any DrawingML color child under parent_el into (rgb, alpha_val)."""
    if parent_el is None:
        return None
    for ch in parent_el:
        ctag = ch.tag.split("}")[-1]
        base = None
        if ctag == "srgbClr" and ch.get("val"):
            base = hex_to_rgb(ch.get("val"))
        elif ctag == "schemeClr" and ch.get("val"):
            sval = ch.get("val")
            if sval == "phClr" and ph_rgb is not None:
                base = ph_rgb
            else:
                base = theme_map.get(sval, (0, 0, 0))
        elif ctag == "sysClr":
            base = hex_to_rgb(ch.get("lastClr") or "000000")
        elif ctag == "prstClr" and ch.get("val"):
            base = PRESET_COLORS.get(ch.get("val").lower(), (0, 0, 0))
        elif ctag == "scrgbClr":
            try:
                r = int(round(int(ch.get("r", "0")) / 100000.0 * 255))
                g = int(round(int(ch.get("g", "0")) / 100000.0 * 255))
                b = int(round(int(ch.get("b", "0")) / 100000.0 * 255))
                base = (r, g, b)
            except ValueError:
                base = (0, 0, 0)
        if base is not None:
            return apply_oxml_color_mods(base, ch)
    return None


def write_srgb_into_parent(parent_el, hex_val, alpha_val=None):
    """Replaces any color node under parent_el with explicit <a:srgbClr val='HEX'>."""
    if parent_el is None:
        return
    color_tags = {"srgbClr", "schemeClr", "sysClr", "prstClr", "scrgbClr", "hslClr"}
    for ch in list(parent_el):
        if ch.tag.split("}")[-1] in color_tags:
            parent_el.remove(ch)
    srgb = etree.SubElement(parent_el, qn_a("srgbClr"))
    srgb.set("val", hex_val.upper())
    if alpha_val is not None:
        a_el = etree.SubElement(srgb, qn_a("alpha"))
        a_el.set("val", str(alpha_val))


class DeckColorMapper:
    """
    Maps all colors from an arbitrary source presentation into the official
    IfG / Arsipa Corporate Design palette while strictly preserving luminance
    tiers (dark / medium / light) and distinguishing multiple accent colors.
    """
    def __init__(self, user_prs):
        self.dark_clusters = []   # [(center_rgb, mapped_hex)]
        self.mid_clusters = []
        self.light_clusters = []
        self._scan_presentation(user_prs)

    def _scan_presentation(self, user_prs):
        dark_counts = Counter()
        mid_counts = Counter()
        light_counts = Counter()

        for slide in user_prs.slides:
            tmap = extract_slide_theme_map(slide)
            root = slide._element
            for parent in root.iter():
                ptag = parent.tag.split("}")[-1]
                if ptag in ("solidFill", "gs", "fillRef", "lnRef"):
                    # Skip drop shadows
                    anc_tags = {a.tag.split("}")[-1] for a in parent.iterancestors()}
                    if "effectLst" in anc_tags:
                        continue
                    res = resolve_color_from_parent(parent, tmap)
                    if res is None:
                        continue
                    rgb, _ = res
                    sat = rgb_saturation(rgb)
                    lum = perceived_brightness(rgb)
                    if sat < 0.12:
                        continue
                    # Quantize slightly to group identical brand colors
                    q = (rgb[0] // 8 * 8, rgb[1] // 8 * 8, rgb[2] // 8 * 8)
                    if lum < 0.38:
                        dark_counts[q] += 1
                    elif lum < 0.76:
                        mid_counts[q] += 1
                    else:
                        light_counts[q] += 1

        self.dark_clusters = self._assign_clusters(dark_counts, DARK_BRAND_SWATCHES)
        self.mid_clusters = self._assign_clusters(mid_counts, MID_BRAND_SWATCHES)
        self.light_clusters = self._assign_clusters(light_counts, LIGHT_BRAND_SWATCHES)

    @staticmethod
    def _assign_clusters(counter, swatches):
        clusters = []
        for q_rgb, _ in counter.most_common():
            merged = False
            for existing_rgb, _ in clusters:
                if rgb_distance(q_rgb, existing_rgb) < 34:
                    merged = True
                    break
            if not merged:
                swatch = swatches[len(clusters) % len(swatches)]
                clusters.append((q_rgb, swatch))
        return clusters

    def _match_cluster(self, rgb, clusters, swatches):
        for c_rgb, hex_val in clusters:
            if rgb_distance(rgb, c_rgb) < 40:
                return hex_val
        idx = len(clusters) % len(swatches)
        swatch = swatches[idx]
        clusters.append((rgb, swatch))
        return swatch

    def map_fill_or_line(self, rgb, role="fill"):
        sat = rgb_saturation(rgb)
        lum = perceived_brightness(rgb)

        # Neutral / Achromatic colors
        if sat < 0.12:
            if lum >= 0.94:
                return IFG_WHITE
            elif lum >= 0.86:
                return "F1F5F9"
            elif lum >= 0.70:
                return "CBD5E1"
            elif lum >= 0.32:
                return IFG_MUTED
            else:
                return IFG_NAVY if role == "fill" else IFG_CHARCOAL

        # Chromatic colors mapped by luminance tier so contrast never inverts
        if lum < 0.38:
            return self._match_cluster(rgb, self.dark_clusters, DARK_BRAND_SWATCHES)
        elif lum < 0.76:
            return self._match_cluster(rgb, self.mid_clusters, MID_BRAND_SWATCHES)
        else:
            return self._match_cluster(rgb, self.light_clusters, LIGHT_BRAND_SWATCHES)


# ----------------- 2. FONT & TEXT RESTYLING HELPERS -----------------

def is_symbol_text(text):
    if not text:
        return False
    stripped = text.strip()
    if not stripped:
        return False
    # Private Use Area glyphs (\uF000-\uF8FF) used by Wingdings/Symbol in OpenXML
    return any(0xF000 <= ord(ch) <= 0xF8FF for ch in stripped)


def is_symbol_font_name(typeface):
    if not typeface:
        return False
    return typeface.strip().lower() in SYMBOL_FONTS


def apply_corporate_font_to_rpr(rpr_el, run_text=""):
    """Sets Abadi on <a:rPr> / <a:defRPr> / <a:endParaRPr> while protecting symbol glyphs/fonts."""
    if rpr_el is None or is_symbol_text(run_text):
        return

    latin = rpr_el.find(qn_a("latin"))
    if latin is not None and is_symbol_font_name(latin.get("typeface")):
        return

    if latin is None:
        # Insert <a:latin> in proper DrawingML child order (after Fill/Effect/Highlight, before ea/cs/sym)
        latin = etree.Element(qn_a("latin"))
        insert_after_tags = {
            "ln", "noFill", "solidFill", "gradFill", "blipFill", "pattFill", "grpFill",
            "effectLst", "effectDag", "highlight", "uLnTx", "uLn", "uFillTx", "uFill"
        }
        idx = 0
        for i, ch in enumerate(rpr_el):
            if ch.tag.split("}")[-1] in insert_after_tags:
                idx = i + 1
        rpr_el.insert(idx, latin)

    latin.attrib.clear()
    latin.set("typeface", CORPORATE_FONT)

    for tag_name in ("cs", "ea"):
        f_el = rpr_el.find(qn_a(tag_name))
        if f_el is not None and not is_symbol_font_name(f_el.get("typeface")):
            f_el.attrib.clear()
            f_el.set("typeface", CORPORATE_FONT)


def restyle_txbody(txbody, theme_map, mapper, bg_lum, is_heading_shape=False, style_font_rgb=None, font_scale=1.0):
    """
    Restyles all runs in a <p:txBody> / <a:txBody> to Abadi + Corporate Colors
    while preserving 100% of original paragraphs, line breaks, spacing, bullets,
    alignments, bold/italic, and ensuring readable contrast against bg_lum.
    """
    if txbody is None:
        return

    # Restyle lstStyle defRPr fonts if present
    lst_style = txbody.find(qn_a("lstStyle"))
    lst_lvl_colors = {}
    lst_lvl_sizes = {}
    if lst_style is not None:
        for ch in lst_style:
            local = ch.tag.split("}")[-1]
            if local.startswith("lvl") and local.endswith("pPr"):
                lvl_num = local[3]
                def_rpr = ch.find(qn_a("defRPr"))
                if def_rpr is not None:
                    res = resolve_color_from_parent(def_rpr.find(qn_a("solidFill")), theme_map)
                    if res is not None:
                        lst_lvl_colors[lvl_num] = res[0]
                    if def_rpr.get("sz"):
                        try:
                            lst_lvl_sizes[lvl_num] = int(def_rpr.get("sz"))
                            if abs(font_scale - 1.0) > 0.02:
                                def_rpr.set("sz", str(int(round(lst_lvl_sizes[lvl_num] * font_scale))))
                        except ValueError:
                            pass
                    def_sf = def_rpr.find(qn_a("solidFill"))
                    if def_sf is not None:
                        write_srgb_into_parent(def_sf, IFG_WHITE if bg_lum < 0.48 else (IFG_NAVY if is_heading_shape else IFG_CHARCOAL))
                    apply_corporate_font_to_rpr(def_rpr)

    # Determine largest font size in this text body to distinguish main title vs kicker line
    max_sz = max(lst_lvl_sizes.values(), default=0)
    for r in txbody.findall(f".//{{{A_NS}}}r"):
        rpr = r.find(qn_a("rPr"))
        if rpr is not None and rpr.get("sz"):
            try:
                max_sz = max(max_sz, int(rpr.get("sz")))
            except ValueError:
                pass

    default_fallback_rgb = style_font_rgb if style_font_rgb is not None else theme_map.get("tx1", (0, 0, 0))

    for p in txbody.findall(qn_a("p")):
        ppr = p.find(qn_a("pPr"))
        lvl_key = str(int(ppr.get("lvl", "0")) + 1) if ppr is not None else "1"
        p_def_rgb = lst_lvl_colors.get(lvl_key, default_fallback_rgb)
        p_def_sz = lst_lvl_sizes.get(lvl_key, 0)

        if ppr is not None:
            def_rpr = ppr.find(qn_a("defRPr"))
            if def_rpr is not None:
                res = resolve_color_from_parent(def_rpr.find(qn_a("solidFill")), theme_map)
                if res is not None:
                    p_def_rgb = res[0]
                if def_rpr.get("sz"):
                    try:
                        p_def_sz = int(def_rpr.get("sz"))
                    except ValueError:
                        pass
                def_sf = def_rpr.find(qn_a("solidFill"))
                if def_sf is not None:
                    write_srgb_into_parent(def_sf, IFG_WHITE if bg_lum < 0.48 else (IFG_NAVY if is_heading_shape else IFG_CHARCOAL))
                apply_corporate_font_to_rpr(def_rpr)

            # Restyle explicit bullet color (<a:buClr>) if present; never touch <a:buFont>!
            bu_clr = ppr.find(qn_a("buClr"))
            if bu_clr is not None:
                new_bu_hex = IFG_WHITE if bg_lum < 0.48 else IFG_NAVY
                write_srgb_into_parent(bu_clr, new_bu_hex)

        for child in p:
            ctag = child.tag.split("}")[-1]
            if ctag == "r":
                t_el = child.find(qn_a("t"))
                r_txt = t_el.text if t_el is not None and t_el.text else ""
                rpr = child.find(qn_a("rPr"))
                if rpr is None:
                    rpr = etree.Element(qn_a("rPr"))
                    child.insert(0, rpr)

                run_sz = p_def_sz
                if rpr.get("sz"):
                    try:
                        run_sz = int(rpr.get("sz"))
                    except ValueError:
                        pass

                # Scale font size proportionally only if slide scale != 1.0
                if abs(font_scale - 1.0) > 0.02 and rpr.get("sz"):
                    try:
                        rpr.set("sz", str(int(round(int(rpr.get("sz")) * font_scale))))
                    except ValueError:
                        pass

                # Resolve original text color
                orig_rgb = p_def_rgb
                orig_sf = rpr.find(qn_a("solidFill"))
                if orig_sf is not None:
                    res = resolve_color_from_parent(orig_sf, theme_map)
                    if res is not None:
                        orig_rgb = res[0]

                # Choose corporate text color with strict contrast guarantee
                if bg_lum < 0.48:
                    new_txt_hex = IFG_WHITE
                else:
                    orig_sat = rgb_saturation(orig_rgb)
                    orig_lum = perceived_brightness(orig_rgb)
                    if orig_sat >= 0.20 and orig_lum < 0.78:
                        # Highlighted / colored text run -> map to dark/readable corporate brand color
                        mapped = mapper.map_fill_or_line(orig_rgb, role="fill")
                        if perceived_brightness(hex_to_rgb(mapped)) > 0.55:
                            mapped = "006174"
                        new_txt_hex = mapped
                    elif is_heading_shape:
                        if max_sz > 0 and run_sz > 0 and run_sz < max_sz - 300 and run_sz < 1800:
                            new_txt_hex = IFG_MUTED
                        else:
                            new_txt_hex = IFG_NAVY
                    elif run_sz >= 2000:
                        new_txt_hex = IFG_NAVY
                    elif 0.40 <= orig_lum <= 0.78 and orig_sat < 0.20 and (run_sz == 0 or run_sz < 1800):
                        # Small muted gray kicker / caption
                        new_txt_hex = IFG_MUTED
                    else:
                        new_txt_hex = IFG_CHARCOAL

                # Ensure <a:solidFill> exists in proper DrawingML schema position inside <a:rPr>
                for ftag in ("noFill", "solidFill", "gradFill", "pattFill"):
                    old_f = rpr.find(qn_a(ftag))
                    if old_f is not None:
                        rpr.remove(old_f)

                sf = etree.Element(qn_a("solidFill"))
                write_srgb_into_parent(sf, new_txt_hex)
                insert_pos = 1 if (len(rpr) > 0 and rpr[0].tag.split("}")[-1] == "ln") else 0
                rpr.insert(insert_pos, sf)

                apply_corporate_font_to_rpr(rpr, r_txt)

            elif ctag == "fld":
                rpr = child.find(qn_a("rPr"))
                if rpr is not None:
                    fld_sf = rpr.find(qn_a("solidFill"))
                    if fld_sf is not None:
                        write_srgb_into_parent(fld_sf, IFG_WHITE if bg_lum < 0.48 else IFG_CHARCOAL)
                    apply_corporate_font_to_rpr(rpr)
            elif ctag == "endParaRPr":
                end_sf = child.find(qn_a("solidFill"))
                if end_sf is not None:
                    write_srgb_into_parent(end_sf, IFG_WHITE if bg_lum < 0.48 else (IFG_NAVY if is_heading_shape else IFG_CHARCOAL))
                apply_corporate_font_to_rpr(child)


# ----------------- 3. PLACEHOLDER & GEOMETRY MATERIALIZATION -----------------

def find_inherited_placeholder(old_slide, ph_idx, ph_type):
    """Finds matching layout and master placeholder elements for style inheritance."""
    matches = []
    for container in (old_slide.slide_layout, old_slide.slide_layout.slide_master):
        try:
            for sh in container.placeholders:
                if sh.placeholder_format.idx == ph_idx or (ph_type is not None and sh.placeholder_format.type == ph_type):
                    matches.append(sh._element)
                    break
        except Exception:
            pass
    return matches


def extract_inherited_level_props(ph_elems, is_title_ph=False):
    """Extracts inherited {lvl_int: {'sz': str, 'b': str, 'algn': str, 'buNone': bool}} from layout/master."""
    props = {i: {} for i in range(9)}
    # Walk master first, then layout so layout overrides master
    for el in reversed(ph_elems):
        lst = el.find(f".//{{{P_NS}}}txBody/{{{A_NS}}}lstStyle")
        if lst is None:
            continue
        for lvl_idx in range(9):
            lp = lst.find(qn_a(f"lvl{lvl_idx + 1}pPr"))
            if lp is None:
                continue
            if lp.get("algn"):
                props[lvl_idx]["algn"] = lp.get("algn")
            if lp.find(qn_a("buNone")) is not None:
                props[lvl_idx]["buNone"] = True
            elif any(lp.find(qn_a(bt)) is not None for bt in ("buChar", "buAutoNum", "buBlip")):
                props[lvl_idx]["buNone"] = False
            dr = lp.find(qn_a("defRPr"))
            if dr is not None:
                if dr.get("sz"):
                    props[lvl_idx]["sz"] = dr.get("sz")
                if dr.get("b"):
                    props[lvl_idx]["b"] = dr.get("b")

    if is_title_ph:
        if "sz" not in props[0]:
            props[0]["sz"] = "2800"
        props[0].setdefault("buNone", True)
    return props


def should_skip_source_shape(sh):
    """Skips hidden 0x0 think-cell metadata and empty/system placeholders (slide number, date, empty footer)."""
    name_low = (sh.name or "").lower()
    if "think-cell" in name_low:
        return True

    cNvPr = sh._element.find(f".//{{{P_NS}}}cNvPr")
    if cNvPr is not None and cNvPr.get("hidden") == "1":
        w = sh.width or 0
        h = sh.height or 0
        if w <= 5000 and h <= 5000:
            return True

    if sh.is_placeholder:
        try:
            pht = sh.placeholder_format.type
            if pht in (PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.HEADER, PP_PLACEHOLDER.FOOTER):
                return True
        except Exception:
            pass

        tag = sh._element.tag.split("}")[-1]
        if tag == "sp":
            has_txt = getattr(sh, "has_text_frame", False) and sh.text_frame and sh.text_frame.text.strip()
            has_fill = sh._element.find(f"{{{P_NS}}}spPr/{{{A_NS}}}blipFill") is not None
            if not has_txt and not has_fill:
                return True

    if any(k in name_low for k in ("slide number", "foliennummer", "page number", "datum")):
        if not (getattr(sh, "has_text_frame", False) and len(sh.text_frame.text.strip()) > 12):
            return True

    return False


def insert_sppr_child(sp_pr, new_el, slot_rank):
    """
    Inserts a DrawingML child into <p:spPr> strictly respecting CT_ShapeProperties sequence:
    0: xfrm | 1: custGeom, prstGeom | 2: fills | 3: ln | 4: effects | 5: scene3d | 6: sp3d | 7: extLst
    """
    rank_map = {
        "xfrm": 0,
        "custGeom": 1, "prstGeom": 1,
        "noFill": 2, "solidFill": 2, "gradFill": 2, "blipFill": 2, "pattFill": 2, "grpFill": 2,
        "ln": 3,
        "effectLst": 4, "effectDag": 4,
        "scene3d": 5, "sp3d": 6, "extLst": 7,
    }
    pos = len(sp_pr)
    for i, ch in enumerate(sp_pr):
        if rank_map.get(ch.tag.split("}")[-1], 99) > slot_rank:
            pos = i
            break
    sp_pr.insert(pos, new_el)


def insert_ppr_bunone(ppr):
    """Inserts <a:buNone/> into <a:pPr> strictly before tabLst/defRPr/extLst."""
    after_bullet_tags = {"tabLst", "defRPr", "extLst"}
    pos = len(ppr)
    for i, ch in enumerate(ppr):
        if ch.tag.split("}")[-1] in after_bullet_tags:
            pos = i
            break
    ppr.insert(pos, etree.Element(qn_a("buNone")))


def ensure_explicit_xfrm(elem, x, y, cx, cy):
    """Ensures a shape/picture/graphicFrame XML element has explicit coordinates."""
    tag = elem.tag.split("}")[-1]
    if tag == "graphicFrame":
        xfrm = elem.find(qn_p("xfrm"))
        if xfrm is None:
            xfrm = etree.Element(qn_p("xfrm"))
            elem.insert(1, xfrm)
    elif tag == "grpSp":
        grp_pr = elem.find(qn_p("grpSpPr"))
        if grp_pr is None:
            grp_pr = etree.Element(qn_p("grpSpPr"))
            elem.insert(1, grp_pr)
        xfrm = grp_pr.find(qn_a("xfrm"))
        if xfrm is None:
            xfrm = etree.Element(qn_a("xfrm"))
            grp_pr.insert(0, xfrm)
    else:
        sp_pr = elem.find(qn_p("spPr"))
        if sp_pr is None:
            sp_pr = etree.Element(qn_p("spPr"))
            elem.insert(1, sp_pr)
        xfrm = sp_pr.find(qn_a("xfrm"))
        if xfrm is None:
            xfrm = etree.Element(qn_a("xfrm"))
            sp_pr.insert(0, xfrm)
        if tag == "pic" and sp_pr.find(qn_a("prstGeom")) is None and sp_pr.find(qn_a("custGeom")) is None:
            prst = etree.Element(qn_a("prstGeom"))
            prst.set("prst", "rect")
            etree.SubElement(prst, qn_a("avLst"))
            insert_sppr_child(sp_pr, prst, slot_rank=1)

    off = xfrm.find(qn_a("off"))
    if off is None:
        off = etree.Element(qn_a("off"))
        xfrm.insert(0, off)
    off.set("x", str(int(round(x))))
    off.set("y", str(int(round(y))))

    ext = xfrm.find(qn_a("ext"))
    if ext is None:
        ext = etree.Element(qn_a("ext"))
        xfrm.insert(1, ext)
    ext.set("cx", str(max(0, int(round(cx)))))
    ext.set("cy", str(max(0, int(round(cy)))))


def materialize_placeholder_and_transform(sh, elem, old_slide, scale, dx, dy, is_top_left_header=False, max_header_right=None):
    """
    Materializes inherited placeholder coordinates/typography if sh.is_placeholder,
    strips <p:ph> so it never collides with Vorlage placeholders, and applies
    aspect-ratio-safe slide coordinate mapping.
    """
    orig_x = sh.left if sh.left is not None else 0
    orig_y = sh.top if sh.top is not None else 0
    orig_cx = sh.width if sh.width is not None else 0
    orig_cy = sh.height if sh.height is not None else 0

    if sh.is_placeholder:
        ph_idx = None
        ph_type = None
        is_title_ph = False
        try:
            ph_idx = sh.placeholder_format.idx
            ph_type = sh.placeholder_format.type
            is_title_ph = ph_type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
        except Exception:
            pass

        inherited_els = find_inherited_placeholder(old_slide, ph_idx, ph_type)
        lvl_props = extract_inherited_level_props(inherited_els, is_title_ph=is_title_ph)

        txbody = elem.find(qn_p("txBody"))
        if txbody is not None:
            for p in txbody.findall(qn_a("p")):
                ppr = p.find(qn_a("pPr"))
                lvl = int(ppr.get("lvl", "0")) if ppr is not None else 0
                inh = lvl_props.get(min(lvl, 8), {})

                if inh.get("algn") or inh.get("buNone"):
                    if ppr is None:
                        ppr = etree.Element(qn_a("pPr"))
                        p.insert(0, ppr)
                    if inh.get("algn") and not ppr.get("algn"):
                        ppr.set("algn", inh["algn"])
                    if inh.get("buNone") and not any(
                        ppr.find(qn_a(bt)) is not None for bt in ("buNone", "buChar", "buAutoNum", "buBlip")
                    ):
                        insert_ppr_bunone(ppr)

                for r in p.findall(qn_a("r")):
                    rpr = r.find(qn_a("rPr"))
                    if rpr is None:
                        rpr = etree.Element(qn_a("rPr"))
                        r.insert(0, rpr)
                    if inh.get("sz") and not rpr.get("sz"):
                        rpr.set("sz", inh["sz"])
                    if inh.get("b") and not rpr.get("b"):
                        rpr.set("b", inh["b"])

        # Remove <p:ph> from <p:nvPr> so it stands as an exact-coordinate native object
        for nv_pr in elem.findall(f".//{{{P_NS}}}nvPr"):
            for ph in list(nv_pr.findall(qn_p("ph"))):
                nv_pr.remove(ph)

    # Compute target coordinates
    new_y = orig_y * scale + dy
    new_cy = orig_cy * scale
    new_cx = orig_cx * scale

    if is_top_left_header:
        # Align top-left header neatly next to Vorlage's left brand bar (0.47") and clear of top-right logo
        new_x = max(int(0.47 * 914400), int(round(orig_x * scale)))
        if max_header_right is not None and (new_x + new_cx) > max_header_right:
            new_cx = max(int(2.0 * 914400), max_header_right - new_x)
    else:
        new_x = orig_x * scale + dx

    ensure_explicit_xfrm(elem, new_x, new_y, new_cx, new_cy)

    # Scale table column widths and row heights only if uniform scale != 1.0
    if abs(scale - 1.0) > 0.02 and elem.tag.split("}")[-1] == "graphicFrame":
        for gc in elem.findall(f".//{{{A_NS}}}gridCol"):
            if gc.get("w"):
                try:
                    gc.set("w", str(int(round(int(gc.get("w")) * scale))))
                except ValueError:
                    pass
        for tr in elem.findall(f".//{{{A_NS}}}tr"):
            if tr.get("h"):
                try:
                    tr.set("h", str(int(round(int(tr.get("h")) * scale))))
                except ValueError:
                    pass

    return (new_x, new_y, new_x + new_cx, new_y + new_cy)


# ----------------- 4. DEEP SHAPE, GROUP, TABLE, SMARTART & CHART RESTYLER -----------------

def get_bg_lum_under_box(bbox, z_bg_stack, default_lum=1.0):
    """Checks previously drawn opaque shapes in Z-order to find the effective background luminance."""
    if not bbox:
        return default_lum
    cx = (bbox[0] + bbox[2]) / 2.0
    cy = (bbox[1] + bbox[3]) / 2.0
    for bx1, by1, bx2, by2, blum in reversed(z_bg_stack):
        if bx1 <= cx <= bx2 and by1 <= cy <= by2:
            return blum
    return default_lum


def get_local_bbox(el):
    """Reads (x1, y1, x2, y2) from an element's local xfrm."""
    xfrm = el.find(f"{{{P_NS}}}spPr/{{{A_NS}}}xfrm")
    if xfrm is None:
        xfrm = el.find(f"{{{P_NS}}}grpSpPr/{{{A_NS}}}xfrm")
    if xfrm is None:
        xfrm = el.find(qn_p("xfrm"))
    if xfrm is None:
        return None
    off = xfrm.find(qn_a("off"))
    ext = xfrm.find(qn_a("ext"))
    if off is None or ext is None:
        return None
    try:
        x = int(off.get("x", "0"))
        y = int(off.get("y", "0"))
        cx = int(ext.get("cx", "0"))
        cy = int(ext.get("cy", "0"))
        return (x, y, x + cx, y + cy)
    except ValueError:
        return None


def restyle_sp_or_cxn(el, theme_map, mapper, z_bg_stack, default_bg_lum=1.0, is_heading_shape=False, font_scale=1.0):
    """Restyles a <p:sp> or <p:cxnSp> (fills, gradients, borders, style refs, and text runs)."""
    tag = el.tag.split("}")[-1]
    sp_pr = el.find(qn_p("spPr"))
    style_el = el.find(qn_p("style"))
    bbox = get_local_bbox(el)
    underlying_lum = get_bg_lum_under_box(bbox, z_bg_stack, default_bg_lum)

    effective_bg_lum = underlying_lum
    has_opaque_fill = False

    # 1. Restyle Fill (<p:spPr> or <p:style>/<a:fillRef>)
    if sp_pr is not None:
        no_fill = sp_pr.find(qn_a("noFill"))
        solid_fill = sp_pr.find(qn_a("solidFill"))
        grad_fill = sp_pr.find(qn_a("gradFill"))
        blip_fill = sp_pr.find(qn_a("blipFill"))

        if no_fill is not None:
            effective_bg_lum = underlying_lum
        elif blip_fill is not None:
            has_opaque_fill = True
            effective_bg_lum = 0.5
        elif solid_fill is not None:
            res = resolve_color_from_parent(solid_fill, theme_map)
            if res is not None:
                rgb, alpha_val = res
                new_hex = mapper.map_fill_or_line(rgb, role="fill")
                write_srgb_into_parent(solid_fill, new_hex, alpha_val)
                fill_lum = perceived_brightness(hex_to_rgb(new_hex))
                alpha_f = (int(alpha_val) / 100000.0) if alpha_val is not None else 1.0
                effective_bg_lum = fill_lum * alpha_f + underlying_lum * (1.0 - alpha_f)
                if alpha_f > 0.25:
                    has_opaque_fill = True
        elif grad_fill is not None:
            gs_list = grad_fill.findall(f".//{{{A_NS}}}gs")
            stop_hexes = []
            for gs in gs_list:
                res = resolve_color_from_parent(gs, theme_map)
                if res is not None:
                    rgb, alpha_val = res
                    new_hex = mapper.map_fill_or_line(rgb, role="fill")
                    if stop_hexes and new_hex == stop_hexes[-1]:
                        new_hex = shift_hex_lightness(new_hex, 0.10 if perceived_brightness(rgb) > 0.3 else 0.14)
                    stop_hexes.append(new_hex)
                    write_srgb_into_parent(gs, new_hex, alpha_val)
            if stop_hexes:
                avg_lum = sum(perceived_brightness(hex_to_rgb(h)) for h in stop_hexes) / len(stop_hexes)
                effective_bg_lum = avg_lum
                has_opaque_fill = True
        elif style_el is not None:
            fill_ref = style_el.find(qn_a("fillRef"))
            if fill_ref is not None and fill_ref.get("idx", "0") != "0":
                res = resolve_color_from_parent(fill_ref, theme_map)
                if res is not None:
                    rgb, alpha_val = res
                    new_hex = mapper.map_fill_or_line(rgb, role="fill")
                    write_srgb_into_parent(fill_ref, new_hex, alpha_val)
                    # Also materialize explicit solidFill so Vorlage fmtScheme doesn't alter brightness
                    sf = etree.Element(qn_a("solidFill"))
                    write_srgb_into_parent(sf, new_hex, alpha_val)
                    insert_sppr_child(sp_pr, sf, slot_rank=2)
                    effective_bg_lum = perceived_brightness(hex_to_rgb(new_hex))
                    has_opaque_fill = True

    # 2. Restyle Border / Connector Line (<a:ln> and <p:style>/<a:lnRef>)
    ln_el = sp_pr.find(qn_a("ln")) if sp_pr is not None else None
    if ln_el is not None:
        ln_sf = ln_el.find(qn_a("solidFill"))
        if ln_sf is not None:
            res = resolve_color_from_parent(ln_sf, theme_map)
            if res is not None:
                rgb, alpha_val = res
                new_ln_hex = mapper.map_fill_or_line(rgb, role="line")
                write_srgb_into_parent(ln_sf, new_ln_hex, alpha_val)

    if style_el is not None:
        ln_ref = style_el.find(qn_a("lnRef"))
        if ln_ref is not None and ln_ref.get("idx", "0") != "0":
            res = resolve_color_from_parent(ln_ref, theme_map)
            if res is not None:
                rgb, alpha_val = res
                new_ln_hex = mapper.map_fill_or_line(rgb, role="line")
                write_srgb_into_parent(ln_ref, new_ln_hex, alpha_val)

    # 3. Resolve <p:style>/<a:fontRef> default text color & restyle <p:txBody>
    style_font_rgb = None
    if style_el is not None:
        font_ref = style_el.find(qn_a("fontRef"))
        if font_ref is not None:
            res = resolve_color_from_parent(font_ref, theme_map)
            if res is not None:
                style_font_rgb = res[0]
            write_srgb_into_parent(font_ref, IFG_WHITE if effective_bg_lum < 0.48 else IFG_CHARCOAL)

    if tag == "sp":
        txbody = el.find(qn_p("txBody"))
        restyle_txbody(
            txbody,
            theme_map=theme_map,
            mapper=mapper,
            bg_lum=effective_bg_lum,
            is_heading_shape=is_heading_shape,
            style_font_rgb=style_font_rgb,
            font_scale=font_scale
        )

    if has_opaque_fill and bbox is not None:
        z_bg_stack.append((bbox[0], bbox[1], bbox[2], bbox[3], effective_bg_lum))


def restyle_table_element(tbl_el, theme_map, mapper, font_scale=1.0):
    """Restyles an existing <a:tbl> in-place (preserving merged cells, column widths, row heights, borders)."""
    tbl_pr = tbl_el.find(qn_a("tblPr"))
    has_first_row_style = tbl_pr is not None and tbl_pr.get("firstRow") == "1" and tbl_pr.find(qn_a("tableStyleId")) is not None

    for r_idx, tr in enumerate(tbl_el.findall(qn_a("tr"))):
        for tc in tr.findall(qn_a("tc")):
            tc_pr = tc.find(qn_a("tcPr"))
            cell_bg_lum = 0.12 if (r_idx == 0 and has_first_row_style) else 1.0

            if tc_pr is not None:
                sf = tc_pr.find(qn_a("solidFill"))
                if sf is not None:
                    res = resolve_color_from_parent(sf, theme_map)
                    if res is not None:
                        rgb, alpha_val = res
                        new_hex = mapper.map_fill_or_line(rgb, role="fill")
                        write_srgb_into_parent(sf, new_hex, alpha_val)
                        cell_bg_lum = perceived_brightness(hex_to_rgb(new_hex))
                elif tc_pr.find(qn_a("noFill")) is not None:
                    cell_bg_lum = 1.0

                # Restyle explicit cell border colors cleanly
                for b_tag in ("lnL", "lnR", "lnT", "lnB"):
                    b_el = tc_pr.find(qn_a(b_tag))
                    if b_el is not None:
                        b_sf = b_el.find(qn_a("solidFill"))
                        if b_sf is not None:
                            res = resolve_color_from_parent(b_sf, theme_map)
                            if res is not None:
                                rgb, alpha_val = res
                                write_srgb_into_parent(b_sf, mapper.map_fill_or_line(rgb, role="line"), alpha_val)

            txbody = tc.find(qn_a("txBody"))
            restyle_txbody(
                txbody,
                theme_map=theme_map,
                mapper=mapper,
                bg_lum=cell_bg_lum,
                is_heading_shape=(r_idx == 0),
                style_font_rgb=(255, 255, 255) if cell_bg_lum < 0.48 else (0, 0, 0),
                font_scale=font_scale
            )


def restyle_any_element(el, theme_map, mapper, z_bg_stack, default_bg_lum=1.0, is_heading_shape=False, font_scale=1.0):
    """Dispatches deep color & font restyling across shapes, connectors, pictures, groups, and tables."""
    tag = el.tag.split("}")[-1]
    if tag in ("sp", "cxnSp"):
        restyle_sp_or_cxn(el, theme_map, mapper, z_bg_stack, default_bg_lum, is_heading_shape, font_scale)
    elif tag == "pic":
        sp_pr = el.find(qn_p("spPr"))
        if sp_pr is not None:
            ln_sf = sp_pr.find(f"{{{A_NS}}}ln/{{{A_NS}}}solidFill")
            if ln_sf is not None:
                res = resolve_color_from_parent(ln_sf, theme_map)
                if res is not None:
                    write_srgb_into_parent(ln_sf, mapper.map_fill_or_line(res[0], role="line"), res[1])
    elif tag == "grpSp":
        grp_stack = []
        for child in el:
            ctag = child.tag.split("}")[-1]
            if ctag in ("sp", "cxnSp", "pic", "grpSp", "graphicFrame"):
                restyle_any_element(child, theme_map, mapper, grp_stack, default_bg_lum, False, font_scale)
    elif tag == "graphicFrame":
        tbl = el.find(f".//{{{A_NS}}}tbl")
        if tbl is not None:
            restyle_table_element(tbl, theme_map, mapper, font_scale=font_scale)


# ----------------- 5. SMARTART, CHART & RELATIONSHIP SAFE CLONER -----------------

def restyle_smartart_or_chart_xml(xml_bytes, partname_str, theme_map, mapper):
    """Updates fonts to Abadi and harmonizes colors inside cloned SmartArt / Chart XML parts."""
    try:
        root = etree.fromstring(xml_bytes)
    except Exception:
        return xml_bytes

    pname = partname_str.lower()
    old_acc1_lum = perceived_brightness(theme_map.get("accent1", (68, 114, 196)))

    # 1. Update all text fonts to Abadi (skipping symbol fonts)
    for rpr in root.findall(f".//{{{A_NS}}}rPr") + root.findall(f".//{{{A_NS}}}defRPr") + root.findall(f".//{{{A_NS}}}endParaRPr"):
        parent = rpr.getparent()
        t_el = parent.find(qn_a("t")) if parent is not None else None
        r_txt = t_el.text if t_el is not None and t_el.text else ""
        apply_corporate_font_to_rpr(rpr, r_txt)

    # 2. If SmartArt colors.xml or drawing.xml relied on a light/medium accent1 in source theme,
    # map accent1 -> accent2 (#3EB9C0 Corporate Cyan/Teal in Vorlage) so shapes stay vibrant & readable!
    if "/diagrams/" in pname and old_acc1_lum >= 0.45:
        for sc in root.findall(f".//{{{A_NS}}}schemeClr"):
            if sc.get("val") == "accent1":
                sc.set("val", "accent2")

    # 3. Map explicit <a:srgbClr> colors (excluding shadow effectLst) into corporate palette
    for parent in root.iter():
        ptag = parent.tag.split("}")[-1]
        if ptag in ("solidFill", "gs"):
            anc_tags = {a.tag.split("}")[-1] for a in parent.iterancestors()}
            if "effectLst" in anc_tags:
                continue
            srgb = parent.find(qn_a("srgbClr"))
            if srgb is not None and srgb.get("val"):
                res = resolve_color_from_parent(parent, theme_map)
                if res is not None:
                    rgb, alpha_val = res
                    # In rPr (text run), keep dark/white readability
                    if "rPr" in anc_tags or "defRPr" in anc_tags:
                        lum = perceived_brightness(rgb)
                        new_hex = IFG_WHITE if lum > 0.85 else IFG_NAVY
                    else:
                        new_hex = mapper.map_fill_or_line(rgb, role="fill")
                    write_srgb_into_parent(parent, new_hex, alpha_val)

    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def allocate_unique_packuri(orig_uri_str, used_partnames):
    """Generates a guaranteed unique PackURI in destination package."""
    if orig_uri_str not in used_partnames:
        used_partnames.add(orig_uri_str)
        return PackURI(orig_uri_str)

    base, ext = posixpath.splitext(orig_uri_str)
    m = re.match(r"^(.*?)(\d+)$", base)
    prefix = m.group(1) if m else base + "_"
    counter = int(m.group(2)) + 100 if m else 101
    while True:
        candidate = f"{prefix}{counter}{ext}"
        if candidate not in used_partnames:
            used_partnames.add(candidate)
            return PackURI(candidate)
        counter += 1


def clone_part_into_package(old_part, dst_pkg, used_partnames, part_memo, theme_map, mapper, extra_xml_transform=None):
    """Recursively clones an OPC Part (image, SVG, SmartArt, chart, workbook, media) without name collisions."""
    key = id(old_part)
    if key in part_memo and extra_xml_transform is None:
        return part_memo[key]

    ctype = getattr(old_part, "content_type", "") or ""
    blob = getattr(old_part, "blob", b"")

    # Raster images -> deduplicated native ImagePart via SHA1
    if ctype.startswith("image/") and "svg" not in ctype and "emf" not in ctype and "wmf" not in ctype and "wdp" not in ctype:
        try:
            img_part = dst_pkg.get_or_add_image_part(io.BytesIO(blob))
            used_partnames.add(str(img_part.partname))
            part_memo[key] = img_part
            return img_part
        except Exception:
            pass

    orig_uri = str(old_part.partname)
    new_uri = allocate_unique_packuri(orig_uri, used_partnames)

    new_blob = blob
    if ctype.endswith("+xml") and ("/diagrams/" in orig_uri or "/charts/" in orig_uri):
        new_blob = restyle_smartart_or_chart_xml(new_blob, orig_uri, theme_map, mapper)
    if extra_xml_transform is not None:
        new_blob = extra_xml_transform(new_blob)

    new_part = Part(new_uri, ctype, new_blob, dst_pkg)
    if extra_xml_transform is None:
        part_memo[key] = new_part

    # Recursively clone sub-relationships (e.g. embedded Excel in chart, images in SmartArt)
    sub_rid_map = {}
    for sub_rid, sub_rel in old_part.rels.items():
        if sub_rel.is_external:
            nr = new_part.relate_to(sub_rel.target_ref, sub_rel.reltype, is_external=True)
        else:
            cloned_sub = clone_part_into_package(
                sub_rel.target_part, dst_pkg, used_partnames, part_memo, theme_map, mapper
            )
            nr = new_part.relate_to(cloned_sub, sub_rel.reltype)
        if nr != sub_rid:
            sub_rid_map[sub_rid] = nr

    if sub_rid_map and ctype.endswith("+xml"):
        try:
            sub_root = etree.fromstring(new_part._blob)
            for node in sub_root.iter():
                for attr_k, attr_v in list(node.attrib.items()):
                    if attr_v in sub_rid_map:
                        node.set(attr_k, sub_rid_map[attr_v])
            new_part._blob = etree.tostring(sub_root, xml_declaration=True, encoding="UTF-8", standalone=True)
        except Exception:
            pass

    return new_part


def copy_and_restyle_shape(sh, elem, old_slide, new_slide, theme_map, mapper, used_partnames, part_memo):
    """
    Transfers all rId relationships of elem from old_slide to new_slide without
    any partname or rId collisions, including SmartArt diagramDrawing parts,
    and assigns unique cNvPr IDs on new_slide.
    """
    dst_pkg = new_slide.part.package

    # Collect all rId values referenced in attributes of elem
    referenced_rids = set()
    for node in elem.iter():
        for attr_v in node.attrib.values():
            if re.match(r"^rId\d+$", attr_v):
                referenced_rids.add(attr_v)

    # SmartArt special handling: data.xml contains <dsp:dataModelExt relId="rIdX"/>
    # referencing a diagramDrawing rel on the SLIDE!
    rid_map = {}
    for rid in sorted(referenced_rids):
        if rid not in old_slide.part.rels:
            continue
        old_rel = old_slide.part.rels[rid]
        if old_rel.is_external:
            new_rid = new_slide.part.relate_to(old_rel.target_ref, old_rel.reltype, is_external=True)
            rid_map[rid] = new_rid
        elif old_rel.reltype.endswith("/diagramData"):
            # Inspect data.xml for dsp:dataModelExt relId
            data_blob = old_rel.target_part.blob
            ext_m = re.search(rb'relId="(rId\d+)"', data_blob)
            draw_rid_new = None
            draw_rid_old = None
            if ext_m:
                draw_rid_old = ext_m.group(1).decode("ascii")
                if draw_rid_old in old_slide.part.rels:
                    d_rel = old_slide.part.rels[draw_rid_old]
                    cloned_draw = clone_part_into_package(
                        d_rel.target_part, dst_pkg, used_partnames, {}, theme_map, mapper
                    )
                    draw_rid_new = new_slide.part.relate_to(cloned_draw, d_rel.reltype)

            def _patch_data_relid(b, old_r=draw_rid_old, new_r=draw_rid_new):
                if old_r and new_r and old_r != new_r:
                    return b.replace(f'relId="{old_r}"'.encode("ascii"), f'relId="{new_r}"'.encode("ascii"))
                return b

            cloned_data = clone_part_into_package(
                old_rel.target_part, dst_pkg, used_partnames, {}, theme_map, mapper,
                extra_xml_transform=_patch_data_relid
            )
            new_rid = new_slide.part.relate_to(cloned_data, old_rel.reltype)
            rid_map[rid] = new_rid
        else:
            # For diagram layout/quickStyle/colors, clone per-diagram instance cleanly
            memo_to_use = {} if "/diagrams/" in str(getattr(old_rel.target_part, "partname", "")) else part_memo
            cloned_target = clone_part_into_package(
                old_rel.target_part, dst_pkg, used_partnames, memo_to_use, theme_map, mapper
            )
            new_rid = new_slide.part.relate_to(cloned_target, old_rel.reltype)
            rid_map[rid] = new_rid

    # Rewrite rId attributes in-place on elem
    if rid_map:
        for node in elem.iter():
            for attr_k, attr_v in list(node.attrib.items()):
                if attr_v in rid_map:
                    node.set(attr_k, rid_map[attr_v])

    # Assign unique cNvPr IDs across the slide
    existing_ids = set()
    for c in new_slide.shapes._spTree.iter(qn_p("cNvPr")):
        try:
            existing_ids.add(int(c.get("id")))
        except (TypeError, ValueError):
            pass
    next_id = max(existing_ids | {1}) + 1
    for c in elem.iter(qn_p("cNvPr")):
        c.set("id", str(next_id))
        next_id += 1

    new_slide.shapes._spTree.append(elem)


def sanitize_pptx_package(buf, slide_count):
    """
    Cleans template leftovers that PowerPoint Desktop (Windows 11) and SharePoint Online refuse:
    - Deduplicates all ZIP entry filenames (prevents corrupt package warnings)
    - Removes ppt/changesInfos/* co-authoring history pointing at deleted template slides
    - Updates docProps/app.xml slide counts
    - Strips orphaned content-type overrides and broken relationships
    """
    import zipfile
    buf.seek(0)
    src = zipfile.ZipFile(buf)

    seen_names = set()
    deduped_names = []
    for n in src.namelist():
        if n not in seen_names:
            seen_names.add(n)
            deduped_names.append(n)

    drop = {
        n for n in deduped_names
        if n.startswith("ppt/changesInfos/") or n.startswith("ppt/revisionInfo") or n.startswith("[trash]/")
    }
    keep = [n for n in deduped_names if n not in drop]
    keep_set = set(keep)

    app_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
        'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
        '<TotalTime>0</TotalTime><Application>Microsoft Office PowerPoint</Application>'
        '<PresentationFormat>Breitbild</PresentationFormat>'
        f'<Slides>{slide_count}</Slides><Notes>0</Notes><HiddenSlides>0</HiddenSlides><MMClips>0</MMClips>'
        '<ScaleCrop>false</ScaleCrop><LinksUpToDate>false</LinksUpToDate><SharedDoc>false</SharedDoc>'
        '<HyperlinksChanged>false</HyperlinksChanged><AppVersion>16.0000</AppVersion></Properties>'
    )

    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        written_entries = set()
        order = ["[Content_Types].xml"] + [n for n in keep if n != "[Content_Types].xml"]
        for n in order:
            if n in written_entries:
                continue
            written_entries.add(n)
            data = src.read(n)
            if n == "[Content_Types].xml":
                xml = data.decode("utf-8")
                seen_overrides = set()

                def _ok(m):
                    pname = m.group(1)
                    if pname not in keep_set or pname in seen_overrides:
                        return ""
                    seen_overrides.add(pname)
                    return m.group(0)

                xml = re.sub(r'<Override PartName="/([^"]+)"[^>]*/>', _ok, xml)
                data = xml.encode("utf-8")
            elif n == "docProps/app.xml":
                data = app_xml.encode("utf-8")
            elif n.endswith(".rels"):
                xml = data.decode("utf-8")
                base_dir = posixpath.dirname(posixpath.dirname(n))

                def _rel_ok(m):
                    tag = m.group(0)
                    if 'TargetMode="External"' in tag:
                        return tag
                    t = re.search(r'Target="([^"]+)"', tag).group(1)
                    full = posixpath.normpath(posixpath.join(base_dir, t)).lstrip("/") if not t.startswith("/") else t.lstrip("/")
                    return tag if full in keep_set else ""

                xml = re.sub(r"<Relationship [^>]*/>", _rel_ok, xml)
                data = xml.encode("utf-8")
            dst.writestr(n, data)
    out.seek(0)
    return out


# ----------------- 6. 1:1 EXACT-LAYOUT CORPORATE TRANSFORMATION ENGINE -----------------

def detect_slide_heading_shape(shapes):
    """Finds the primary top heading/title shape on a slide (if any) without altering layout."""
    best = None
    best_score = -999
    for sh in shapes:
        if should_skip_source_shape(sh):
            continue
        if not (getattr(sh, "has_text_frame", False) and sh.text_frame and sh.text_frame.text.strip()):
            continue
        txt = sh.text_frame.text.strip()
        if len(txt) > 180:
            continue
        top_in = (sh.top or 0) / 914400.0
        left_in = (sh.left or 0) / 914400.0
        score = 0
        if top_in < 1.15:
            score += 45
        elif top_in < 2.2:
            score += 15
        else:
            score -= 25

        if left_in < 2.0:
            score += 10
        if len(txt) < 90:
            score += 20

        max_pt = 0
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if r.font.size:
                    max_pt = max(max_pt, r.font.size.pt)
        if max_pt >= 20:
            score += 25

        if sh.is_placeholder:
            try:
                if sh.placeholder_format.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
                    score += 50
            except Exception:
                pass

        if score > best_score and score >= 25:
            best_score = score
            best = sh
    return best


def build_arsipa_presentation(template_file, user_file):
    """
    Converts every slide (N -> N) from user_file into template_file (Vorlage),
    preserving 100% of the exact original layout, shapes, groups, tables, images,
    SmartArt diagrams, and charts while restyling ONLY fonts (Abadi) and colors
    (IfG / Arsipa Corporate Design).
    """
    output_prs = Presentation(template_file)
    user_prs = Presentation(user_file)

    catalog = scan_arsipa_template(output_prs)
    canvas_layout = output_prs.slide_layouts[catalog["canvas_leere_folie"]]

    # Pre-scan deck colors for consistent 1:1 corporate color mapping
    mapper = DeckColorMapper(user_prs)

    # Clear template sample slides safely (Single Presentation Instance rule)
    slide_ids = list(output_prs.slides._sldIdLst)
    for sld_id in slide_ids:
        r_id = sld_id.rId
        output_prs.part.drop_rel(r_id)
        output_prs.slides._sldIdLst.remove(sld_id)

    # Strip stale template section lists referencing deleted slide IDs
    for ext in output_prs.part._element.xpath("./p:extLst/p:ext"):
        if ext.xpath('./*[local-name()="sectionLst"]'):
            ext.getparent().remove(ext)

    # Compute aspect-ratio-safe coordinate mapping from user_prs -> Vorlage
    old_w = float(user_prs.slide_width)
    old_h = float(user_prs.slide_height)
    new_w = float(output_prs.slide_width)
    new_h = float(output_prs.slide_height)

    sx = new_w / old_w
    sy = new_h / old_h
    scale = min(sx, sy)
    dx = (new_w - old_w * scale) / 2.0
    dy = (new_h - old_h * scale) / 2.0
    aspect_differs = abs(sx - sy) > 0.04
    max_header_right = int(11.85 * 914400)  # Keep top headers clear of top-right Vorlage logo

    pkg = output_prs.part.package
    pkg_parts = list(pkg.iter_parts()) if hasattr(pkg, "iter_parts") else getattr(pkg, "parts", [])
    used_partnames = {str(p.partname) for p in pkg_parts}
    part_memo = {}
    processed_slides = []

    for idx, old_slide in enumerate(user_prs.slides):
        new_slide = output_prs.slides.add_slide(canvas_layout)

        # Remove empty cloned layout placeholders on new_slide so no ghost boxes overlap user shapes
        for ph in list(new_slide.placeholders):
            try:
                new_slide.shapes._spTree.remove(ph._element)
            except Exception:
                pass

        theme_map = extract_slide_theme_map(old_slide)
        slide_bg_lum = 1.0

        # If old_slide has an explicit custom slide background (<p:bg>), restyle its color too
        old_bg = old_slide._element.find(f"{{{P_NS}}}cSld/{{{P_NS}}}bg")
        if old_bg is not None:
            new_bg = deepcopy(old_bg)
            bg_sf = new_bg.find(f".//{{{A_NS}}}solidFill")
            if bg_sf is not None:
                res = resolve_color_from_parent(bg_sf, theme_map)
                if res is not None:
                    mapped_bg = mapper.map_fill_or_line(res[0], role="fill")
                    write_srgb_into_parent(bg_sf, mapped_bg, res[1])
                    slide_bg_lum = perceived_brightness(hex_to_rgb(mapped_bg))
                    csld = new_slide._element.find(qn_p("cSld"))
                    existing_bg = csld.find(qn_p("bg"))
                    if existing_bg is not None:
                        csld.remove(existing_bg)
                    csld.insert(0, new_bg)

        source_shapes = list(old_slide.shapes)
        heading_sh = detect_slide_heading_shape(source_shapes)
        slide_title_label = ""
        if heading_sh is not None and getattr(heading_sh, "has_text_frame", False):
            slide_title_label = " ".join(heading_sh.text_frame.text.split())

        z_bg_stack = []
        copied_count = 0

        for sh in source_shapes:
            if should_skip_source_shape(sh):
                continue

            if not slide_title_label and getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip():
                slide_title_label = " ".join(sh.text_frame.text.split())[:60]

            elem = deepcopy(sh._element)

            is_top_left_hdr = (
                aspect_differs
                and sh == heading_sh
                and (sh.top or 0) / 914400.0 < 1.05
                and (sh.left or 0) / 914400.0 < 1.30
            )

            bbox = materialize_placeholder_and_transform(
                sh, elem, old_slide,
                scale=scale, dx=dx, dy=dy,
                is_top_left_header=is_top_left_hdr,
                max_header_right=max_header_right
            )

            restyle_any_element(
                elem,
                theme_map=theme_map,
                mapper=mapper,
                z_bg_stack=z_bg_stack,
                default_bg_lum=slide_bg_lum,
                is_heading_shape=(sh == heading_sh),
                font_scale=scale
            )

            copy_and_restyle_shape(
                sh, elem, old_slide, new_slide,
                theme_map=theme_map,
                mapper=mapper,
                used_partnames=used_partnames,
                part_memo=part_memo
            )
            copied_count += 1

        processed_slides.append({
            "index": idx,
            "title": slide_title_label or f"Folie {idx + 1}",
            "layout_name": canvas_layout.name,
            "reason": f"1:1 Original-Aufbau erhalten ({copied_count} Elemente) • Schriftart Abadi & IfG-Farben angewendet"
        })

    out_buf = io.BytesIO()
    output_prs.save(out_buf)
    out_buf = sanitize_pptx_package(out_buf, len(processed_slides))

    return out_buf, processed_slides


# ----------------- STREAMLIT UI: 1-CLICK CLEAN WORKFLOW -----------------

logo_html = f'<img src="data:image/png;base64,{logo_b64}" alt="IfG Logo" class="ifg-logo-img" />' if logo_b64 else '<strong style="color: #012641; font-size: 1.3rem;">IfG</strong>'

st.markdown(f"""
<div class="ifg-top-nav">
    <div class="ifg-nav-left">
        {logo_html}
        <div class="ifg-nav-divider"></div>
        <div>
            <div class="ifg-nav-title">
                IfG SlideStyler
                <span class="ifg-badge-tool">Internes Tool</span>
            </div>
            <p class="ifg-nav-subtitle">PowerPoint Vorlagen-Engine nach offiziellem Corporate Design</p>
        </div>
    </div>
    <div class="ifg-nav-badge">
        <svg width="14" height="14" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round">
            <rect x="3" y="11" width="18" height="11" rx="2" ry="2"></rect>
            <path d="M7 11V7a5 5 0 0 1 10 0v4"></path>
        </svg>
        <span>100% DSGVO-konform (Lokale Verarbeitung)</span>
    </div>
</div>
""", unsafe_allow_html=True)

st.markdown("""
<div class="ifg-card">
    <div class="ifg-card-header">
        <svg width="20" height="20" viewBox="0 0 24 24" fill="none" stroke="#012641" stroke-width="2.2" stroke-linecap="round" stroke-linejoin="round">
            <polygon points="12 2 15.09 8.26 22 9.27 17 14.14 18.18 21.02 12 17.77 5.82 21.02 7 14.14 2 9.27 8.91 8.26 12 2"></polygon>
        </svg>
        <h3 class="ifg-card-title">1:1 Layout-Erhaltung & Corporate Design (Nur Farben & Schriftart)</h3>
    </div>
    <ul class="ifg-feature-list">
        <li class="ifg-feature-item">
            <div class="ifg-bullet-dot"></div>
            <div><strong>Exakte 1:1 Anordnung</strong>: Alle Grafiken, Formen, Gruppen, Tabellen, Bilder und SmartArt-Diagramme bleiben exakt in ihrer ursprünglichen Struktur und Position erhalten.</div>
        </li>
        <li class="ifg-feature-item">
            <div class="ifg-bullet-dot"></div>
            <div><strong>Offizielle Typografie (<code>Abadi</code>)</strong>: Sämtliche Texte (inkl. Tabellen, Gruppen und Diagramme) werden auf <code>Abadi</code> umgestellt, während Symbole & Icons erhalten bleiben.</div>
        </li>
        <li class="ifg-feature-item">
            <div class="ifg-bullet-dot"></div>
            <div><strong>Intelligente IfG / Arsipa Farbpalette</strong>: Alte Farben werden kontrastgetreu in Deep Navy (<code>#012641</code>), Corporate Teal (<code>#00A8B7</code> / <code>#3EB9C0</code>) und Slate Charcoal (<code>#3A4A63</code>) überführt.</div>
        </li>
        <li class="ifg-feature-item">
            <div class="ifg-bullet-dot"></div>
            <div><strong>100% Editierbar & Kompatibel</strong>: Alle Folien (egal ob 10 oder 50+ Seiten) bleiben vollständig nativ bearbeitbar und öffnen fehlerfrei unter Windows 11 & SharePoint Online.</div>
        </li>
    </ul>
</div>
""", unsafe_allow_html=True)

DEFAULT_TEMPLATES = {
    "⭐ IfG Vorlage Kurzfassung (Empfohlen – 6,7 MB, schnelle Ladezeit)": "Arsipa Vorlage/ifg_PowerPoint_Vorlage_Kurzfassung.pptx",
    "🏢 IfG Vorlage Vollversion (47 Layouts – 37 MB)": "Arsipa Vorlage/Vorlage_Power_Point_Arsipa.pptx",
    "📁 Eigene Vorlagendatei (.pptx) hochladen": None
}

template_choice = st.selectbox(
    "1. Wählen Sie die gewünschte Unternehmensvorlage aus (Template Selection):",
    list(DEFAULT_TEMPLATES.keys()),
    index=0
)

template_file_to_use = None

if DEFAULT_TEMPLATES[template_choice] is not None:
    template_path = DEFAULT_TEMPLATES[template_choice]
    if os.path.exists(template_path):
        template_file_to_use = template_path
        st.success(f"✅ Vorlage „{os.path.basename(template_path)}“ erfolgreich geladen (Einsatzbereit)")
    else:
        st.error(f"Vorlagendatei nicht gefunden unter: {template_path}")
else:
    template_uploaded = st.file_uploader("Laden Sie Ihre eigene Vorlagendatei (.pptx) hoch:", type=["pptx"])
    if template_uploaded:
        template_file_to_use = template_uploaded

user_file = st.file_uploader("2. Ausgangspräsentation hochladen, die konvertiert werden soll (Ihre .pptx-Datei):", type=["pptx"])

if template_file_to_use and user_file:
    st.markdown("<div style='height: 8px;'></div>", unsafe_allow_html=True)

    current_key = f"{getattr(user_file, 'name', '')}_{getattr(user_file, 'size', 0)}_{template_choice}"
    if st.session_state.get('last_processed_file') != current_key:
        st.session_state.pop('transformed_bytes', None)
        st.session_state.pop('transformed_slides', None)
        st.session_state.pop('transformed_filename', None)

    if st.button("🚀 Präsentation jetzt konvertieren (1-Click Transform)"):
        with st.spinner("⚡ Folien werden 1:1 in die Vorlage übertragen (Farben & Schriftarten werden angepasst)..."):
            try:
                out_buf, processed_slides = build_arsipa_presentation(template_file_to_use, user_file)
                st.session_state['transformed_bytes'] = out_buf.getvalue()
                st.session_state['transformed_slides'] = processed_slides
                st.session_state['transformed_filename'] = f"IfG_{user_file.name}"
                st.session_state['last_processed_file'] = current_key
            except Exception as ex:
                import traceback
                tb_str = traceback.format_exc()
                st.error(f"Verarbeitungsfehler aufgetreten: {ex}")
                st.code(tb_str)
                print(tb_str)

    if 'transformed_bytes' in st.session_state:
        processed_slides = st.session_state['transformed_slides']
        out_bytes = st.session_state['transformed_bytes']
        out_filename = st.session_state['transformed_filename']

        st.markdown(f"""
        <div class="ifg-success-box">
            <div class="ifg-success-title">🎉 Konvertierung erfolgreich abgeschlossen!</div>
            <p class="ifg-success-desc">
                Alle <strong>{len(processed_slides)} Folien</strong> wurden im exakten 1:1 Original-Layout in die IfG / Arsipa Vorlage übertragen – mit angepasster <code>Abadi</code>-Typografie und offiziellen Unternehmensfarben.
            </p>
        </div>
        """, unsafe_allow_html=True)

        st.download_button(
            label="📥 Konvertierte Präsentation herunterladen (.pptx)",
            data=out_bytes,
            file_name=out_filename,
            mime="application/vnd.openxmlformats-officedocument.presentationml.presentation"
        )

        with st.expander("🔍 Folien-Übersicht anzeigen"):
            for s in processed_slides:
                st.markdown(f"**Folie {s['index']+1}:** `{s['title'][:55]}` ➔ {s['reason']}")
else:
    st.info("Bitte wählen Sie eine Vorlage aus und laden Sie Ihre Ausgangspräsentation hoch, um die Konvertierung zu starten.")
